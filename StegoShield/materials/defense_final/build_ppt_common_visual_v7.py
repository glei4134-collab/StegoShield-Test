# -*- coding: utf-8 -*-
from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(r"D:\Desktop\StegoShield_Defense_Final_2026-04-17")
ASSET_DIR = ROOT / "assets_common"
OUTPUT = ROOT / "defense_ppt_common_visual_v7.pptx"


def font(size, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simsun.ttc",
    ]
    for p in candidates:
        try:
            return ImageFont.truetype(p, size)
        except Exception:
            continue
    return ImageFont.load_default()


def make_canvas(title):
    img = Image.new("RGB", (1400, 900), (245, 248, 255))
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1400, 90), fill=(25, 85, 180))
    d.text((40, 24), title, fill=(255, 255, 255), font=font(40, True))
    return img, d


def save_chart_01():
    img, d = make_canvas("用户需求占比（调研样本）")
    labels = ["隐蔽传输", "易用操作", "数据恢复", "成本可控"]
    vals = [36, 28, 22, 14]
    colors = [(46, 134, 222), (0, 184, 148), (241, 196, 15), (231, 76, 60)]
    x = 120
    for i, (lb, v) in enumerate(zip(labels, vals)):
        h = int(v * 12)
        y = 780 - h
        d.rectangle((x, y, x + 180, 780), fill=colors[i])
        d.text((x + 30, y - 38), f"{v}%", fill=(44, 62, 80), font=font(30, True))
        d.text((x + 18, 800), lb, fill=(44, 62, 80), font=font(24))
        x += 290
    img.save(ASSET_DIR / "chart_bar.png")


def save_chart_02():
    img, d = make_canvas("项目阶段目标（大二→大三）")
    points = [(150, 700), (450, 560), (760, 460), (1080, 350)]
    for i in range(len(points) - 1):
        d.line((points[i], points[i + 1]), fill=(52, 152, 219), width=8)
    texts = ["MVP完成", "校内试点", "稳定优化", "场景扩展"]
    for (x, y), t in zip(points, texts):
        d.ellipse((x - 20, y - 20, x + 20, y + 20), fill=(41, 128, 185))
        d.text((x - 70, y - 70), t, fill=(44, 62, 80), font=font(24, True))
    img.save(ASSET_DIR / "chart_line.png")


def save_chart_03():
    img, d = make_canvas("技术流程图")
    boxes = [
        (80, 280, 300, 420, "输入文本/文件"),
        (360, 280, 580, 420, "AES加密"),
        (640, 280, 860, 420, "隐写嵌入"),
        (920, 280, 1140, 420, "输出图片"),
    ]
    for b in boxes:
        d.rounded_rectangle(b[:4], radius=18, fill=(230, 240, 255), outline=(52, 152, 219), width=4)
        d.text((b[0] + 20, b[1] + 50), b[4], fill=(44, 62, 80), font=font(26, True))
    for i in range(3):
        x = 300 + i * 280
        d.polygon([(x + 30, 350), (x + 80, 330), (x + 80, 370)], fill=(52, 152, 219))
    img.save(ASSET_DIR / "flow.png")


def save_chart_04():
    img, d = make_canvas("竞品对比（核心维度）")
    x0, y0 = 110, 170
    row_h = 110
    col = [260, 260, 260, 260, 260]
    heads = ["维度", "竞品A", "竞品B", "StegoShield", "结论"]
    rows = [
        ["隐蔽性", "中", "低", "高", "我们更强"],
        ["易用性", "中", "高", "高", "持平领先"],
        ["加密能力", "无", "弱", "AES-256", "明显领先"],
        ["成本", "高", "中", "低", "适合学生场景"],
    ]
    xx = x0
    for i, w in enumerate(col):
        d.rectangle((xx, y0, xx + w, y0 + row_h), fill=(41, 128, 185))
        d.text((xx + 16, y0 + 36), heads[i], fill=(255, 255, 255), font=font(24, True))
        xx += w
    for r, row in enumerate(rows):
        yy = y0 + (r + 1) * row_h
        xx = x0
        for c, w in enumerate(col):
            d.rectangle((xx, yy, xx + w, yy + row_h), outline=(120, 160, 210), width=2, fill=(245, 250, 255))
            d.text((xx + 16, yy + 36), row[c], fill=(44, 62, 80), font=font(23))
            xx += w
    img.save(ASSET_DIR / "table_compare.png")


def save_chart_05():
    img, d = make_canvas("推广漏斗（学生阶段）")
    levels = [
        (180, 180, 1220, 320, "课程展示触达"),
        (260, 330, 1140, 450, "兴趣用户咨询"),
        (350, 460, 1050, 560, "实际试用"),
        (430, 570, 970, 650, "持续复用"),
    ]
    cols = [(52, 152, 219), (46, 204, 113), (241, 196, 15), (231, 76, 60)]
    for i, lv in enumerate(levels):
        d.rounded_rectangle(lv[:4], radius=14, fill=cols[i])
        d.text((lv[0] + 24, lv[1] + 34), lv[4], fill=(255, 255, 255), font=font(28, True))
    img.save(ASSET_DIR / "funnel.png")


def build_assets():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    save_chart_01()
    save_chart_02()
    save_chart_03()
    save_chart_04()
    save_chart_05()


def add_slide_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(11, 21, 40)
    bg.line.fill.background()


def add_slide_title(slide, title, subtitle):
    t = slide.shapes.add_textbox(Inches(0.7), Inches(0.35), Inches(8.4), Inches(1.1))
    tf = t.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 248, 255)

    s = slide.shapes.add_textbox(Inches(0.7), Inches(1.25), Inches(8.8), Inches(0.8))
    sf = s.text_frame
    sf.clear()
    q = sf.paragraphs[0]
    q.text = subtitle
    q.font.name = "Microsoft YaHei"
    q.font.size = Pt(20)
    q.font.color.rgb = RGBColor(158, 190, 240)


def add_points(slide, lines):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.7), Inches(2.05), Inches(6.4), Inches(4.95))
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(24, 40, 70)
    box.fill.transparency = 8
    box.line.color.rgb = RGBColor(80, 140, 230)
    box.line.width = Pt(1.2)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(240, 246, 255)
        p.space_after = Pt(8)


def add_chart(slide, path):
    frame = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.35), Inches(2.05), Inches(5.25), Inches(4.95))
    frame.fill.solid()
    frame.fill.fore_color.rgb = RGBColor(236, 244, 255)
    frame.line.color.rgb = RGBColor(95, 155, 245)
    frame.line.width = Pt(1.1)
    slide.shapes.add_picture(str(path), Inches(7.45), Inches(2.15), width=Inches(5.05), height=Inches(4.75))


def build_ppt():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    chart_map = [
        "flow.png",
        "chart_bar.png",
        "flow.png",
        "flow.png",
        "table_compare.png",
        "chart_bar.png",
        "chart_line.png",
        "funnel.png",
        "chart_line.png",
        "flow.png",
        "chart_bar.png",
        "funnel.png",
        "table_compare.png",
    ]
    content = [
        ("StegoShield 项目答辩", "大二团队：图像隐写 + AES加密 的工程化实践", ["项目目标：做出可运行、可复现、可解释的MVP", "应用场景：校园课程项目与科研协作", "答辩结构：问题-方案-验证-计划", "团队分工：主讲/算法/产品"]),
        ("1. 背景与意义", "为什么我们做这个项目", ["只做加密会暴露通信意图", "隐写让传输过程更隐蔽", "学生用户需要低门槛安全工具", "项目价值：实用且可推广"]),
        ("2. 问题定义", "我们聚焦什么真实问题", ["敏感信息在普通通道传输风险高", "同类工具复杂导致使用率低", "缺乏学生场景可落地方案", "目标：安全、易用、可持续"]),
        ("3. 技术流程", "端到端处理链路", ["输入文本/文件与图片", "可选AES-256加密", "执行隐写嵌入输出载体图", "按密钥提取并恢复原文"]),
        ("4. 工程结构", "不是概念，而是可运行代码", ["后端：Flask API", "模块：隐写/加密/异常处理", "测试：matrix与边界测试脚本", "仓库路径：StegoShield/backend + tests"]),
        ("5. 核心能力", "已完成功能", ["文本与文件嵌入提取", "AES-256 加密保护", "错误码与响应规范", "可重复演示与调试"]),
        ("6. 测试与验证", "用结果证明可行性", ["建立样本测试集", "对比不同输入场景表现", "验证提取稳定性", "持续优化异常处理"]),
        ("7. 商业与推广", "学生阶段的现实策略", ["先校内验证再外拓", "基础免费 + 增值能力", "课程展示与路演获客", "看激活率和留存率"]),
        ("8. 阶段计划", "大二到大三的路线", ["2026：完成MVP与试点", "2027：稳定性与体验优化", "2028：探索校外合作", "每阶段都有可量化指标"]),
        ("9. 成员一（主讲）", "统筹与落地", ["负责路线设计与进度推进", "组织联调与答辩材料", "坚持真实可复现表达", "交接成员二算法部分"]),
        ("10. 成员二（算法）", "实验与鲁棒性", ["负责参数实验与结果复测", "关注提取成功率", "沉淀测试与问题定位流程", "交接成员三产品部分"]),
        ("11. 成员三（产品）", "交互与调研", ["优化嵌入-提取操作体验", "整理访谈反馈优先级", "提升可理解与可用性", "完成商业文档与收束"]),
        ("12. 总结", "我们交付了什么", ["完成了可运行的StegoShield原型", "完成了可展示的测试验证链路", "明确了下一阶段优化方向", "请老师批评指正"]),
    ]

    for i in range(13):
        slide = prs.slides.add_slide(layout)
        add_slide_bg(slide)
        add_slide_title(slide, content[i][0], content[i][1])
        add_points(slide, content[i][2])
        add_chart(slide, ASSET_DIR / chart_map[i])

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build_assets()
    build_ppt()
