# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(r"D:\Desktop\StegoShield_Defense_Final_2026-04-17")
ASSETS = ROOT / "assets_web"
OUTPUT = ROOT / "defense_ppt_visual_pro_v5.pptx"


def add_bg_image(slide, image_path):
    slide.shapes.add_picture(str(image_path), 0, 0, width=Inches(13.333), height=Inches(7.5))
    shade = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    shade.fill.solid()
    shade.fill.fore_color.rgb = RGBColor(8, 14, 26)
    shade.fill.transparency = 28
    shade.line.fill.background()


def add_title(slide, text, y=0.45, size=40):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.8), Inches(1.0))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 248, 255)


def add_subtitle(slide, text, y=1.45, size=20):
    box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.8), Inches(0.8))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(size)
    p.font.color.rgb = RGBColor(194, 211, 255)


def add_bullets(slide, lines):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(8.2), Inches(4.9))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(12, 24, 45)
    panel.fill.transparency = 20
    panel.line.color.rgb = RGBColor(70, 127, 220)
    panel.line.width = Pt(1.2)

    tf = panel.text_frame
    tf.clear()
    tf.word_wrap = True
    for idx, line in enumerate(lines):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.level = 0
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(22)
        p.font.color.rgb = RGBColor(240, 246, 255)
        p.space_after = Pt(8)


def add_right_image(slide, image_path):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.15), Inches(2.0), Inches(3.35), Inches(4.9))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(16, 29, 54)
    card.line.color.rgb = RGBColor(86, 151, 255)
    card.line.width = Pt(1.2)
    slide.shapes.add_picture(str(image_path), Inches(9.25), Inches(2.1), width=Inches(3.15), height=Inches(4.7))


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    imgs = sorted(ASSETS.glob("*.jpg"))
    if len(imgs) < 8:
        raise RuntimeError("图片素材不足，至少需要8张。")

    slides_data = [
        ("StegoShield", "大二团队：图像隐写 + AES加密 的可落地安全项目", ["专业创新创业训练答辩", "成员一（主讲）/成员二（算法）/成员三（产品）", "指导老师：[请填写姓名]", "2026.04.24"]),
        ("1. 选题背景与意义", "为什么我们要做 StegoShield", ["只做加密会暴露通信意图", "隐写让信息“看起来像普通图片”", "个人隐私与校园协作都需要低门槛安全工具", "项目目标：先做可运行，再做可规模化"]),
        ("2. 小组分工", "三人协作，职责清晰", ["成员一：统筹进度、主讲、联调", "成员二：算法实验、鲁棒性测试", "成员三：产品流程、调研与商业文档", "周例会+看板+复盘，保证可追踪产出"]),
        ("3. 项目结构", "从问题到落地的完整链路", ["问题定义：真实场景中的通信隐蔽需求", "技术实现：Flask + 隐写服务 + AES", "产品层：简化嵌入/提取流程", "商业层：校园验证 + 分层策略"]),
        ("4. 解决痛点", "不只“安全”，更要“可用”", ["痛点1：敏感信息直接传输风险高", "痛点2：加密文件容易引起注意", "痛点3：传统工具门槛高，用户不愿长期用", "方案：隐写 + 加密 + 清晰提示"]),
        ("5. 产品优势", "工程化能力 + 真实项目证据", ["支持文本/文件嵌入提取", "支持 AES-256 加密保护", "后端接口、前端页面、测试脚本均已落地", "适配大二团队：能讲清、能演示、能迭代"]),
        ("6. 商业与推广", "先验证，再扩展", ["基础功能免费，高级能力增值", "首批目标用户：校内科研/课程项目组", "推广路径：课程展示、竞赛路演、内容传播", "核心指标：激活率、留存率、复用率"]),
        ("7. 资金与战略", "保守预算，稳步推进", ["阶段策略：2026做MVP试点", "预算重点：服务器、测试、展示与调研", "不追求虚高融资，优先证明可持续性", "2027-2028再扩展校外场景"]),
        ("8. 调研与验证", "我们用数据说话", ["技术调研：LSB/DCT方案对比", "用户调研：可用性与信任感访谈", "竞品调研：易用性/安全性差异分析", "结论：用户最在意“能稳定取回数据”"]),
        ("9. 成员一（主讲）", "负责统筹与落地", ["推进项目闭环：从想法到可演示", "组织协作：任务拆解、节奏控制", "答辩视角：真实问题优先于宏大叙事", "交接：下面请成员二讲算法部分"]),
        ("10. 成员二（算法）", "负责实验与鲁棒性", ["构建样本，做参数对比与复测", "关注点：不止能嵌入，更要稳提取", "坚持测试驱动，结果可复现", "交接：下面请成员三讲产品部分"]),
        ("11. 成员三（产品）", "负责交互与用户调研", ["优化嵌入-保存-提取流程", "改进提示文案，降低使用门槛", "构建反馈闭环，支持持续迭代", "结尾：三位成员汇报完毕"]),
        ("12. 总结", "大二团队也能做出可靠作品", ["我们完成了可运行、可复现、可迭代的MVP", "后续重点：稳定性提升 + 场景拓展", "请老师批评指正，欢迎提问"]),
    ]

    for idx, (title, subtitle, lines) in enumerate(slides_data):
        slide = prs.slides.add_slide(layout)
        bg = imgs[idx % len(imgs)]
        side = imgs[(idx + 2) % len(imgs)]
        add_bg_image(slide, bg)
        add_title(slide, title)
        add_subtitle(slide, subtitle)
        add_bullets(slide, lines)
        add_right_image(slide, side)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
