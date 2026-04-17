from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation


def find_source_files():
    wx_dir = Path(r"C:\Users\17544\Documents\WXWork\1688854959646674\Cache\File\2026-04")
    source_ppt = None
    source_doc = None
    for file_path in wx_dir.glob("*.pptx"):
        if "答辩" in file_path.name:
            source_ppt = file_path
            break
    for file_path in wx_dir.glob("*.docx"):
        if "数学与应用数学" in file_path.name and "答辩记录表" in file_path.name:
            source_doc = file_path
            break
    if not source_ppt or not source_doc:
        raise FileNotFoundError("未找到原始PPT或答辩记录表模板。")
    return source_ppt, source_doc


def set_textbox_lines(shape, lines):
    text_frame = shape.text_frame
    text_frame.clear()
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.text = line


def get_content():
    slide_payload = {
        1: (
            "专业创新创业训练答辩",
            [
                "专业：数学与应用数学",
                "成员：成员一（主讲）/成员二（算法）/成员三（产品）",
                "学号：2023XXXX01 / 2023XXXX02 / 2023XXXX03",
                "指导老师：[请填写老师姓名]",
                "答辩日期：2026.04.24（按现场通知调整）",
            ],
            "StegoShield：基于图像隐写与AES加密的信息保护项目",
        ),
        2: (
            "一、项目选题背景、社会意义",
            [
                "背景：日常通信中“内容加密”常见，但“通信意图暴露”仍是痛点。",
                "问题：仅使用普通加密文件容易引起注意，不利于敏感信息低可见传输。",
                "方案：StegoShield将密文嵌入普通图片，实现“看起来普通、内容受保护”。",
                "社会意义：服务个人隐私保护、校园科研协作、小型团队安全沟通。",
                "大二团队定位：先把技术做实、把场景跑通，再逐步扩展商业化。",
            ],
            None,
        ),
        3: (
            "二、小组分工",
            [
                "成员一（主讲/组长）：项目统筹、接口联调、答辩材料整合。",
                "成员二（算法方向）：DCT/LSB隐写实验、鲁棒性对比、检测指标分析。",
                "成员三（产品方向）：前端交互、用户流程设计、调研与商业模型整理。",
                "协作机制：每周例会+任务看板+代码评审，保证三人贡献可追踪。",
            ],
            None,
        ),
        4: (
            "三、项目计划书内容结构",
            [
                "1）问题与需求：为什么需要“隐写+加密”的双层保护。",
                "2）技术实现：Flask后端 + 图像隐写服务 + AES加密模块。",
                "3）产品设计：嵌入/提取流程、错误提示、可用性优化。",
                "4）商业与推广：校园切入、分层定价、低成本获客。",
                "5）阶段目标：本学期完成MVP验证，下学期推进试点。",
            ],
            None,
        ),
        5: (
            "四、项目解决的痛点问题",
            [
                "痛点1：敏感文本直接传输风险高。",
                "痛点2：仅靠传统加密无法隐藏通信意图。",
                "痛点3：初学者工具门槛高，难以稳定使用。",
                "对应方案：图片隐写隐藏通信意图 + AES-256保护内容 + 简化操作流程。",
                "结果目标：在不增加用户负担前提下提升信息传输安全性。",
            ],
            None,
        ),
        6: (
            "五、项目产品与服务的特色、竞争优势",
            [
                "技术栈真实落地：Flask / numpy / Pillow / pycryptodome / scipy。",
                "核心能力：文本与文件嵌入提取、AES加密、统一错误码返回。",
                "项目证据：已包含后端API、前端页面、测试脚本与实验目录。",
                "竞争优势：教学与实战结合，适配学生团队“可理解、可部署、可演示”。",
                "边界说明：当前以PNG/BMP效果更稳，JPEG场景持续优化中。",
            ],
            None,
        ),
        7: (
            "六、项目商业模式与营销策略",
            [
                "商业模式：基础功能免费 + 高级功能订阅（容量、批处理、团队协作）。",
                "首批用户：校内科研小组、课程项目组、安全兴趣社群。",
                "推广策略：课程展示、比赛路演、技术社区内容发布。",
                "运营原则：先验证留存和口碑，再扩展付费功能。",
                "大二阶段策略：控制成本，优先积累真实使用反馈。",
            ],
            None,
        ),
        8: (
            "七、项目投资与融资资金使用概况",
            [
                "现阶段不追求大额融资，先以“低成本自驱迭代”为主。",
                "资金需求主要用于：云服务器、域名、测试数据与展示物料。",
                "预算建议：5000-10000元完成学年内试点验证。",
                "若后续进入校外试点，再考虑申请校创基金/天使支持。",
                "资金使用原则：技术优先、节奏可控、每笔投入可度量。",
            ],
            None,
        ),
        9: (
            "八、项目财务预测与发展战略",
            [
                "2026（大二）：完成MVP与校内试点，目标100-300种子用户。",
                "2027（大三）：打磨稳定性与易用性，形成轻量付费版本。",
                "2028（大四）：尝试校外合作场景，验证可持续商业路径。",
                "关键指标：激活率、留存率、付费转化、单用户服务成本。",
                "战略核心：技术可信 + 场景可复用 + 团队执行力。",
            ],
            None,
        ),
        10: (
            "九、项目调研活动",
            [
                "代码层调研：对比LSB、DCT及压缩后可提取能力。",
                "用户层调研：围绕“是否会用、是否敢用、是否愿意持续用”访谈。",
                "竞品调研：分析同类工具在安全性、易用性、价格上的差异。",
                "结果：用户最重视“操作简单+明确提示+可恢复数据”。",
                "后续：补充实验样本，完善不同图像质量下的鲁棒性数据。",
                "过渡：下面进入个人部分，由成员一先介绍整体执行。",
            ],
            None,
        ),
        11: (
            "十、成员一（主讲）",
            [
                "工作：统筹项目路线，完成接口联调、部署脚本和整体答辩框架。",
                "见解：大学生项目要先做“能跑且可复现”的最小闭环，不追求虚高叙事。",
                "创业关注点：场景真实、成本可控、节奏稳定、数据透明。",
                "个人反思：以大二能力边界为前提，用工程方法持续迭代。",
                "结尾交接：以上是我的汇报，下面请成员二介绍算法工作。",
            ],
            None,
        ),
        12: (
            "十一、成员二（算法方向）",
            [
                "工作：负责隐写算法实验、鲁棒性测试、指标对比与异常样本分析。",
                "见解：安全方向不能只看“能嵌入”，必须关注“能稳定提取”。",
                "创业关注点：技术路线要与用户价值绑定，避免“为算法而算法”。",
                "个人反思：通过测试驱动，提升了建模与问题定位能力。",
                "结尾交接：我的部分汇报完毕，下面请成员三介绍产品与调研。",
            ],
            None,
        ),
        13: (
            "十二、成员三（产品方向）",
            [
                "工作：负责前端交互、用户流程、问卷访谈与商业文档整理。",
                "见解：真正决定留存的是“首次使用体验”，不是功能数量。",
                "创业关注点：合规意识、用户信任、反馈闭环和长期服务能力。",
                "个人反思：把技术语言翻译成用户语言，是产品落地关键。",
                "结尾收束：三位成员汇报完毕，请老师批评指正。",
            ],
            None,
        ),
    }

    record_text = (
        "【小组必答题1】\n"
        "StegoShield项目的选题背景来自课程实践中的真实观察：仅靠传统加密虽然能保护内容，但会暴露通信意图。"
        "我们希望做一个更贴近学生使用习惯的工具，把敏感信息隐藏在普通图片中，再通过AES加密实现双重保护。"
        "项目社会意义体现在：提升个人隐私保护能力、服务校园科研协作、推动安全知识工程化实践。"
        "计划书结构分为选题与意义、技术实现、产品设计、商业与推广、阶段目标与风险控制。"
        "小组分工为：成员一负责统筹和主讲，成员二负责算法实验与测试，成员三负责产品交互和调研。\n\n"
        "【小组必答题2】\n"
        "本项目痛点聚焦为三点：敏感信息传输风险高、加密通信意图明显、工具门槛较高。"
        "我们以“隐写+加密+可用性”为主线：后端采用Flask构建API，结合numpy、Pillow与pycryptodome完成图像处理和AES-256加密；"
        "工程上提供统一错误码与可复现测试；体验上降低操作门槛。商业模式采用基础免费+增值订阅，先服务校内场景，再逐步扩展。"
        "考虑大二团队阶段特点，现阶段不追求大额融资，主要投入服务器、测试与调研，预算控制在小额可承受范围。"
        "发展战略为2026完成校内MVP验证，2027提升稳定性与留存，2028探索校外合作。"
        "调研活动包括技术对比实验、用户访谈和竞品分析，结论是用户最重视“能稳定取回数据”和“流程足够简单”。\n\n"
        "【个人必答题-成员一（主讲）】\n"
        "我主要负责项目统筹、关键节点推进和答辩组织。通过任务拆解、周例会和联调验证，保障项目从想法到可演示的闭环落地。"
        "我认为大学生创新创业应先解决真实小问题，再逐步放大，不应被宏大叙事牵引。对于大二团队，最重要的是做出可复现、可验证、可解释的成果。"
        "我重点关注四件事：问题真实性、执行可持续、指标可衡量、团队可协同。项目推进中我们也遇到过算法与体验的取舍、样本不足和时间冲突，"
        "我通过优先级管理先保核心闭环，再做性能优化。这个项目提升了我在组织协同、工程推进和压力管理上的能力，也让我更理解创新与落地之间的平衡。\n\n"
        "【个人必答题-成员二（算法方向）】\n"
        "我负责隐写算法实验、鲁棒性测试和结果分析。主要工作是构建测试样本、比较参数在压缩缩放后的可提取率，并把测试结果反馈给产品和后端同学，推动系统优化。"
        "我的核心见解是：安全项目不能只看“能嵌入”，必须看“能稳定提取”。因此我坚持测试驱动，每次改动都进行统一样本复测，保证结果可比较、可复现。"
        "我认为大学生创业应关注技术与价值绑定，避免为了复杂而复杂，同时重视实验日志和复盘机制。"
        "通过这个项目，我在数据分析、问题定位和跨角色沟通能力上有明显提升，也更清楚算法工作的价值最终要体现在用户体验和稳定性上。\n\n"
        "【个人必答题-成员三（产品方向）】\n"
        "我负责产品交互、用户流程和调研文档。我的工作是把技术能力变成可用流程：优化“嵌入-保存-提取”路径，改进文案和错误提示，减少首次使用门槛；"
        "同时组织同学试用访谈，整理需求优先级，参与商业和推广方案编写。我的见解是，决定留存的不是功能数量，而是首次体验是否顺畅、可信和高效。"
        "我认为大学生创业需要关注用户信任、反馈闭环和节奏管理三点：明确数据边界，快速响应问题，避免早期目标过散。"
        "这个项目让我学会把技术语言翻译成用户语言，也让我理解了产品在团队中的桥梁作用。"
    )

    return slide_payload, record_text


def apply_ppt_content(prs, slide_payload):
    for index, slide in enumerate(prs.slides, start=1):
        title, lines, subtitle = slide_payload[index]
        if slide.shapes.title is not None:
            slide.shapes.title.text = title

        if index == 1:
            if len(slide.shapes) > 1 and getattr(slide.shapes[1], "has_text_frame", False):
                slide.shapes[1].text_frame.text = subtitle or ""
            if len(slide.shapes) > 2 and getattr(slide.shapes[2], "has_text_frame", False):
                set_textbox_lines(slide.shapes[2], lines)
            continue

        body_shape = None
        for shape in slide.shapes:
            if shape is not slide.shapes.title and getattr(shape, "has_text_frame", False):
                body_shape = shape
                break
        if body_shape:
            set_textbox_lines(body_shape, lines)


def lock_ppt_fonts(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Microsoft YaHei"


def lock_doc_fonts(doc):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            run.font.name = "宋体"
            if run._element.rPr is not None and run._element.rPr.rFonts is not None:
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.name = "宋体"
                        if run._element.rPr is not None and run._element.rPr.rFonts is not None:
                            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def apply_doc_content(doc, record_text):
    table = doc.tables[0]
    table.cell(0, 2).text = "StegoShield：面向校园场景的图像隐写与加密平台"
    table.cell(1, 2).text = "2023XXXX01"
    table.cell(1, 7).text = "成员一（主讲）"
    table.cell(2, 2).text = "2023XXXX02"
    table.cell(2, 7).text = "成员二（算法）"
    table.cell(3, 2).text = "2023XXXX03"
    table.cell(3, 7).text = "成员三（产品）"
    table.cell(4, 2).text = "待填写"
    table.cell(5, 7).text = "数学楼答辩教室（待确认）"
    table.cell(9, 0).text = record_text


def main():
    out_dir = Path(__file__).resolve().parent
    source_ppt, source_doc = find_source_files()
    slide_payload, record_text = get_content()

    fixed_ppt_path = out_dir / "defense_ppt_final_stegoshield_v3_fixed.pptx"
    fixed_doc_path = out_dir / "defense_record_filled_stegoshield_v3_fixed.docx"
    font_locked_ppt_path = out_dir / "defense_ppt_final_stegoshield_v3_fixed_fontlocked.pptx"
    font_locked_doc_path = out_dir / "defense_record_filled_stegoshield_v3_fixed_fontlocked.docx"

    prs = Presentation(str(source_ppt))
    apply_ppt_content(prs, slide_payload)
    prs.save(str(fixed_ppt_path))

    prs_locked = Presentation(str(source_ppt))
    apply_ppt_content(prs_locked, slide_payload)
    lock_ppt_fonts(prs_locked)
    prs_locked.save(str(font_locked_ppt_path))

    doc = Document(str(source_doc))
    apply_doc_content(doc, record_text)
    doc.save(str(fixed_doc_path))

    doc_locked = Document(str(source_doc))
    apply_doc_content(doc_locked, record_text)
    lock_doc_fonts(doc_locked)
    doc_locked.save(str(font_locked_doc_path))

    print(f"FIXED_PPT={fixed_ppt_path}")
    print(f"FIXED_DOC={fixed_doc_path}")
    print(f"LOCKED_PPT={font_locked_ppt_path}")
    print(f"LOCKED_DOC={font_locked_doc_path}")


if __name__ == "__main__":
    main()
