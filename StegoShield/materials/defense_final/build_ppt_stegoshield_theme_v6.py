# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(r"D:\Desktop\StegoShield_Defense_Final_2026-04-17")
ASSETS = ROOT / "assets_project"
OUTPUT = ROOT / "defense_ppt_stegoshield_theme_v6.pptx"


def add_bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(8, 18, 38)
    bg.line.fill.background()

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.0), Inches(13.333), Inches(0.28))
    band.fill.solid()
    band.fill.fore_color.rgb = RGBColor(0, 208, 255)
    band.line.fill.background()


def add_title(slide, title, subtitle):
    tb = slide.shapes.add_textbox(Inches(0.65), Inches(0.45), Inches(8.8), Inches(1.1))
    tf = tb.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.name = "Microsoft YaHei"
    p.font.size = Pt(38)
    p.font.bold = True
    p.font.color.rgb = RGBColor(245, 250, 255)

    sb = slide.shapes.add_textbox(Inches(0.65), Inches(1.38), Inches(9.2), Inches(0.8))
    sf = sb.text_frame
    sf.clear()
    s = sf.paragraphs[0]
    s.text = subtitle
    s.font.name = "Microsoft YaHei"
    s.font.size = Pt(19)
    s.font.color.rgb = RGBColor(154, 192, 255)


def add_points(slide, lines):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.65), Inches(2.08), Inches(8.65), Inches(4.95))
    panel.fill.solid()
    panel.fill.fore_color.rgb = RGBColor(14, 30, 58)
    panel.fill.transparency = 9
    panel.line.color.rgb = RGBColor(54, 125, 248)
    panel.line.width = Pt(1.25)

    tf = panel.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {line}"
        p.font.name = "Microsoft YaHei"
        p.font.size = Pt(21)
        p.font.color.rgb = RGBColor(242, 247, 255)
        p.space_after = Pt(8)


def add_image_card(slide, img, caption):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.48), Inches(2.08), Inches(3.2), Inches(4.95))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(17, 36, 66)
    card.line.color.rgb = RGBColor(86, 165, 255)
    card.line.width = Pt(1.25)

    slide.shapes.add_picture(str(img), Inches(9.58), Inches(2.18), width=Inches(3.0), height=Inches(3.95))

    cap = slide.shapes.add_textbox(Inches(9.58), Inches(6.18), Inches(3.0), Inches(0.72))
    cf = cap.text_frame
    cf.clear()
    c = cf.paragraphs[0]
    c.text = caption
    c.font.name = "Microsoft YaHei"
    c.font.size = Pt(14)
    c.font.color.rgb = RGBColor(188, 222, 255)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    layout = prs.slide_layouts[6]

    img_order = [
        "最终_E.png",
        "dct_test_original.png",
        "中间_C.png",
        "dct_output.png",
        "test_in.png",
        "en_out.png",
        "rb_out.png",
        "simple_out.png",
        "v2_out.png",
        "test_out_new.png",
        "test_abc.png",
        "test_3char.png",
        "original_A.png",
    ]
    imgs = [ASSETS / name for name in img_order]

    slides = [
        (
            "StegoShield 项目答辩",
            "数学与应用数学（大二）｜图像隐写 + AES 加密 的工程化实践",
            [
                "项目定位：先做可运行、可复现、可解释的最小闭环",
                "技术主线：LSB/DCT实验 + AES加密 + API封装",
                "应用主线：校园场景下的安全信息传递",
                "团队分工：主讲统筹 / 算法实验 / 产品调研",
            ],
            "项目样例图（最终输出）",
        ),
        (
            "1. 选题背景与问题定义",
            "为什么“只加密”不够，为什么要做“隐写 + 加密”",
            [
                "痛点1：敏感内容在普通通道传输风险高",
                "痛点2：仅发送加密文件会暴露通信意图",
                "痛点3：同类工具复杂，普通学生难长期使用",
                "目标：让安全传输既隐蔽又易用",
            ],
            "原始图像样本",
        ),
        (
            "2. 小组分工（大二团队）",
            "三人协作模式：清晰职责 + 周期复盘",
            [
                "成员一：项目统筹、接口联调、答辩框架",
                "成员二：算法测试、鲁棒性对比、问题定位",
                "成员三：交互优化、用户访谈、商业整理",
                "协作方式：每周例会 + 看板 + 可追踪产出",
            ],
            "中间实验图",
        ),
        (
            "3. 技术路线",
            "从输入到提取的可落地流程",
            [
                "输入：文本/文件 + 载体图片",
                "处理：可选AES-256加密，再执行隐写嵌入",
                "输出：视觉近似原图的含密图",
                "提取：按协议读取并可选解密恢复原始信息",
            ],
            "DCT实验输出",
        ),
        (
            "4. 代码与工程结构",
            "项目路径：StegoShield/backend + tests + experiments",
            [
                "后端：Flask API，统一错误码与响应结构",
                "服务：dct_stego/encryption/enhanced_stego/redundancy",
                "测试：test_matrix、边界条件与鲁棒性脚本",
                "优势：不是概念图，而是可运行工程",
            ],
            "输入样例",
        ),
        (
            "5. 核心能力展示",
            "当前已完成能力（以代码仓库为准）",
            [
                "文本/文件嵌入与提取",
                "AES-256 加密保护",
                "错误处理与参数校验",
                "多组样例图验证基本可用性",
            ],
            "加密后输出样例",
        ),
        (
            "6. 验证与测试",
            "我们关注的不只是能嵌入，更是能稳定提取",
            [
                "对比不同样例图的嵌入效果",
                "检查输出图像可读性与可用性",
                "围绕压缩/转换场景做鲁棒性观察",
                "持续完善测试矩阵，降低演示偶发风险",
            ],
            "鲁棒性样例输出",
        ),
        (
            "7. 商业与推广（学生阶段）",
            "先校内验证，再逐步扩展，不讲空泛融资",
            [
                "目标用户：校内课程项目组、科研协作小组",
                "模式：基础免费 + 高级能力增值",
                "推广：课程展示、路演、技术内容传播",
                "核心指标：激活率、留存率、复用率",
            ],
            "简化方案输出",
        ),
        (
            "8. 阶段计划（大二→大三）",
            "用可执行里程碑替代空洞愿景",
            [
                "2026：完成MVP与校内试点验证",
                "2027：优化稳定性、可用性与协作能力",
                "2028：探索校外试点与可持续路径",
                "原则：每一步都能被数据和代码验证",
            ],
            "V2版本输出",
        ),
        (
            "9. 成员一（主讲）",
            "负责统筹推进与落地交付",
            [
                "搭建任务节奏：分工、里程碑、联调",
                "确保答辩材料与系统演示一致",
                "坚持真实可复现，不夸大成果",
                "交接：下面由成员二介绍算法工作",
            ],
            "项目输出样例",
        ),
        (
            "10. 成员二（算法）",
            "负责实验对比与稳定性分析",
            [
                "构建测试样本，做参数对比",
                "重点关注提取稳定性与异常定位",
                "坚持测试驱动，改动后统一复测",
                "交接：下面由成员三介绍产品部分",
            ],
            "算法样例图",
        ),
        (
            "11. 成员三（产品）",
            "负责交互、调研与反馈闭环",
            [
                "优化嵌入-提取流程，降低首次门槛",
                "整理访谈结论，形成需求优先级",
                "将技术语言转成用户可理解表达",
                "结尾：三位成员汇报完毕",
            ],
            "产品侧样例图",
        ),
        (
            "12. 总结与答辩请求",
            "StegoShield：一支大二团队做出的可运行安全项目",
            [
                "我们已完成：可运行、可演示、可复现的MVP",
                "我们在做：稳定性与场景化持续优化",
                "请老师重点指导：算法鲁棒性与应用落地建议",
                "感谢聆听，欢迎提问",
            ],
            "原图参考",
        ),
    ]

    for idx, (title, subtitle, points, caption) in enumerate(slides):
        slide = prs.slides.add_slide(layout)
        add_bg(slide)
        add_title(slide, title, subtitle)
        add_points(slide, points)
        add_image_card(slide, imgs[idx], caption)

    prs.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    build()
